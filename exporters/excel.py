import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List
from models import Company
from config import logger

def export_to_excel(companies: List[Company], file_path: str):
    """
    Exports a list of Company objects to a highly professional, styled Excel file.
    Creates a styled raw data sheet ('Companies') and a dashboard sheet ('Summary').
    """
    if not companies:
        logger.warning("No companies to export.")
        # Create an empty file with headers just in case
        df_empty = pd.DataFrame(columns=[
            "Company Name", "Website", "Country", "City", "Industry", 
            "Startup Category", "Description", "Founding Year", 
            "Employee Count", "Funding Stage", "Total Funding", 
            "LinkedIn", "Email", "Technologies", "Source", "Date Collected"
        ])
        df_empty.to_excel(file_path, index=False)
        return

    logger.info(f"Exporting {len(companies)} companies to Excel: {file_path}")

    # Convert list of Companies to a list of dicts
    data_dicts = [c.to_dict() for c in companies]
    
    # Create DataFrame and rename columns for a professional look
    df = pd.DataFrame(data_dicts)
    column_mapping = {
        "name": "Company Name",
        "website": "Website",
        "country": "Country",
        "city": "City",
        "industry": "Industry",
        "category": "Startup Category",
        "description": "Description",
        "founding_year": "Founding Year",
        "employee_count": "Employee Count",
        "funding_stage": "Funding Stage",
        "total_funding": "Total Funding",
        "linkedin": "LinkedIn",
        "email": "Email",
        "technologies": "Technologies",
        "source": "Source",
        "date_collected": "Date Collected"
    }
    df = df.rename(columns=column_mapping)

    # Reorder columns to match mapping keys' order
    ordered_cols = list(column_mapping.values())
    df = df[ordered_cols]

    # Calculate summaries
    total_companies = len(df)
    
    # Count sources correctly (handling merged comma-separated sources)
    source_counts = {}
    for s in df["Source"]:
        if pd.isna(s):
            continue
        for part in str(s).split(","):
            part_cleaned = part.strip()
            if part_cleaned:
                source_counts[part_cleaned] = source_counts.get(part_cleaned, 0) + 1
    df_sources = pd.DataFrame(list(source_counts.items()), columns=["Source", "Company Count"]).sort_values("Company Count", ascending=False)
    
    # Categories breakdown
    df_categories = df["Startup Category"].value_counts().reset_index()
    df_categories.columns = ["Startup Category", "Company Count"]
    
    # Countries breakdown
    df_countries = df["Country"].value_counts().reset_index()
    df_countries.columns = ["Country", "Company Count"]

    # Write to Excel using openpyxl engine
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        # Write Companies raw sheet
        df.to_excel(writer, sheet_name="Companies", index=False)
        
        # Write Summary sheet placeholder (we will fill manually for premium layout)
        # Create an empty sheet first
        wb = writer.book
        ws_summary = wb.create_sheet("Summary", index=0)
        ws_companies = wb["Companies"]

        # Style the Companies Sheet
        _style_companies_sheet(ws_companies)

        # Style the Summary Sheet
        _build_summary_sheet(ws_summary, total_companies, df_sources, df_categories, df_countries)

    logger.info("Excel export completed successfully.")

def _style_companies_sheet(ws):
    """Styles the 'Companies' raw data sheet."""
    # Colors
    navy_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=11, bold=False)
    
    border_side = Side(border_style="thin", color="D9D9D9")
    border_box = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Enable filters
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    
    # Freeze top row
    ws.freeze_panes = "A2"
    
    # Format Headers (Row 1)
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        
    # Format Data cells
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_body
            cell.border = border_box
            
            # Alignments
            col_name = ws.cell(row=1, column=col_idx).value
            if col_name in ("Founding Year", "Employee Count"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                # Ensure numbers are actual integers if possible
                if cell.value is not None:
                    try:
                        cell.value = int(float(cell.value))
                    except ValueError:
                        pass
            elif col_name in ("Website", "LinkedIn", "Email"):
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single") # Hyperlink style
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Set row heights
    ws.row_dimensions[1].height = 28
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 20

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        
        # Don't make the Description or Technologies column excessively wide
        col_name = col[0].value
        is_long_text = col_name in ("Description", "Technologies")
        
        for cell in col:
            val = str(cell.value or "")
            if is_long_text and len(val) > 40:
                val = val[:40] # cap length for width calculation
            max_len = max(max_len, len(val))
            
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def _build_summary_sheet(ws, total_companies: int, df_sources: pd.DataFrame, df_categories: pd.DataFrame, df_countries: pd.DataFrame):
    """Fills and styles the 'Summary' dashboard sheet."""
    # Hide grid lines? No, keep standard but styled
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_bold_font = Font(name="Calibri", size=11, bold=True)
    body_font = Font(name="Calibri", size=11)
    
    fill_header = PatternFill(start_color="244061", end_color="244061", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    fill_kpi = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Title
    ws["B2"] = "STARTUP FINDER SUMMARY REPORT"
    ws["B2"].font = title_font
    
    # Key Metrics card
    ws.merge_cells("B4:C4")
    ws["B4"] = "KEY METRIC"
    ws["B4"].font = header_font
    ws["B4"].fill = fill_header
    ws["B4"].alignment = Alignment(horizontal="center")
    
    ws.merge_cells("B5:C5")
    ws["B5"] = f"Total Companies: {total_companies}"
    ws["B5"].font = title_font
    ws["B5"].fill = fill_kpi
    ws["B5"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 40
    
    # Add border to KPI card
    for r in range(4, 6):
        for c in range(2, 4):
            ws.cell(row=r, column=c).border = border_thin

    # Helper function to write a summary table
    def write_table(start_row: int, start_col: int, title: str, col_headers: List[str], df_data: pd.DataFrame):
        # Section Title
        ws.cell(row=start_row, column=start_col, value=title).font = section_font
        
        # Headers
        for h_idx, header in enumerate(col_headers):
            cell = ws.cell(row=start_row + 1, column=start_col + h_idx, value=header)
            cell.font = header_font
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center" if h_idx == 1 else "left")
            cell.border = border_thin
            
        # Data
        current_row = start_row + 2
        for r_idx, row in df_data.iterrows():
            # Value
            val_cell = ws.cell(row=current_row, column=start_col, value=row.iloc[0])
            val_cell.font = body_font
            val_cell.border = border_thin
            
            # Count
            count_cell = ws.cell(row=current_row, column=start_col + 1, value=int(row.iloc[1]))
            count_cell.font = body_bold_font
            count_cell.alignment = Alignment(horizontal="right")
            count_cell.border = border_thin
            
            # Zebra striping
            if r_idx % 2 == 1:
                val_cell.fill = fill_zebra
                count_cell.fill = fill_zebra
                
            current_row += 1
            
        # Total Row
        t_label_cell = ws.cell(row=current_row, column=start_col, value="Total")
        t_label_cell.font = body_bold_font
        t_label_cell.border = border_thin
        t_label_cell.fill = fill_kpi
        
        t_val_cell = ws.cell(row=current_row, column=start_col + 1, value=int(df_data.iloc[:, 1].sum()))
        t_val_cell.font = body_bold_font
        t_val_cell.alignment = Alignment(horizontal="right")
        t_val_cell.border = border_thin
        t_val_cell.fill = fill_kpi
        
        # Row heights
        ws.row_dimensions[start_row].height = 22
        ws.row_dimensions[start_row + 1].height = 20
        for r in range(start_row + 2, current_row + 1):
            ws.row_dimensions[r].height = 18

    # Write tables stacked vertically
    # 1. Sources Table
    write_table(7, 2, "Breakdown by Source", ["Source", "Companies"], df_sources)
    
    # 2. Categories Table
    next_start = 7 + len(df_sources) + 5
    write_table(next_start, 2, "Breakdown by Category", ["Category", "Companies"], df_categories)
    
    # 3. Countries Table
    next_start = next_start + len(df_categories) + 5
    write_table(next_start, 2, "Breakdown by Country", ["Country", "Companies"], df_countries)

    # Set column widths for summary sheet
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 15
